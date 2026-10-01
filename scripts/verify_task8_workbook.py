from pathlib import Path

from openpyxl import load_workbook


expected_sheets = {
    "soldiers",
    "duty_shifts",
    "assignments",
    "duty_locations",
    "hierarchy",
    "duty_types",
    "exemption_types",
    "shift_templates",
    "swap_requests",
    "exemption_requests",
    "soldier_field_updates",
    "soldier_enrollment_requests",
    "personal_constraints",
    "soldier_exemptions",
    "system_settings",
    "bug_reports",
    "range_locations",
    "range_events",
    "range_assignments",
    "soldier_range_qualifications",
    "range_excusal_requests",
    "rank_advancement_intervals",
}

path = Path("frontend/tests/e2e/fixtures/storage-import.xlsx")
workbook = load_workbook(path, read_only=True, data_only=True)
assert set(workbook.sheetnames) == expected_sheets, workbook.sheetnames
assert workbook["soldiers"].max_row == 1
assert workbook["soldiers"].max_column == 1
assert workbook["soldiers"]["A1"].value == "personal_number"
print(f"Task 8 workbook verified: {len(workbook.sheetnames)} recognized sheets")
