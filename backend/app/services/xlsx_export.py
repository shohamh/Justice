"""Safe, bounded helpers for creating XLSX files from primitive cell matrices."""

from __future__ import annotations

import math
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict

Cell = str | int | float | bool | None

# The API route enforces this limit on raw request bytes before JSON parsing.
MAX_EXPORT_REQUEST_BYTES = 16 * 1024 * 1024
MAX_EXPORT_ROWS = 25_000
MAX_EXPORT_COLUMNS = 256
MAX_EXPORT_CELLS = 500_000
MAX_EXPORT_CELL_TEXT = 32_767
MAX_EXPORT_SHEETS = 255

_INVALID_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9._ -]+")


class WorksheetMatrix(BaseModel):
    """One worksheet's title, headers, and rows, without scalar coercion."""

    model_config = ConfigDict(strict=True, extra="forbid")

    sheet_name: str
    headers: list[str]
    rows: list[list[Cell]]


def _is_legal_xml_text(value: str) -> bool:
    for character in value:
        codepoint = ord(character)
        if codepoint in (0x9, 0xA, 0xD):
            continue
        if not (
            0x20 <= codepoint <= 0xD7FF
            or 0xE000 <= codepoint <= 0xFFFD
            or 0x10000 <= codepoint <= 0x10FFFF
        ):
            return False
    return True


def _validate_matrices(matrices: list[WorksheetMatrix]) -> None:
    if not matrices:
        raise ValueError("At least one worksheet is required")
    if len(matrices) > MAX_EXPORT_SHEETS:
        raise ValueError(f"Worksheet count exceeds limit of {MAX_EXPORT_SHEETS}")

    total_cells = 0
    titles: set[str] = set()
    total_rows = 0
    for matrix in matrices:
        title = matrix.sheet_name
        if not isinstance(title, str) or not 1 <= len(title) <= 31:
            raise ValueError("Worksheet title must contain 1 to 31 characters")
        if any(character in title for character in "[]:*?/\\"):
            raise ValueError("Worksheet title contains an invalid character")
        if title.startswith("'") or title.endswith("'"):
            raise ValueError("Worksheet title cannot start or end with an apostrophe")
        if not _is_legal_xml_text(title):
            raise ValueError("Worksheet title contains a character that is illegal in XML")
        title_key = title.casefold()
        if title_key in titles:
            raise ValueError("Worksheet titles must be unique")
        titles.add(title_key)

        columns = len(matrix.headers)
        if not 1 <= columns <= MAX_EXPORT_COLUMNS:
            raise ValueError(f"Worksheet column limit is 1 to {MAX_EXPORT_COLUMNS}")
        row_count = len(matrix.rows)
        if row_count > MAX_EXPORT_ROWS:
            raise ValueError(f"Worksheet row limit is {MAX_EXPORT_ROWS}")
        total_rows += row_count
        if total_rows > MAX_EXPORT_ROWS:
            raise ValueError(f"Workbook row limit is {MAX_EXPORT_ROWS}")
        for row in matrix.rows:
            if len(row) != columns:
                raise ValueError("Every data row must match the header width")

        total_cells += columns + row_count * columns
        if total_cells > MAX_EXPORT_CELLS:
            raise ValueError(f"Workbook cell limit is {MAX_EXPORT_CELLS}")

        for text in [
            *matrix.headers,
            *(cell for row in matrix.rows for cell in row if isinstance(cell, str)),
        ]:
            if len(text) > MAX_EXPORT_CELL_TEXT:
                raise ValueError(f"Cell text exceeds length limit of {MAX_EXPORT_CELL_TEXT}")
            if not _is_legal_xml_text(text):
                raise ValueError("Cell text contains a character that is illegal in XML")

        for row in matrix.rows:
            for cell in row:
                if cell is None or type(cell) in (str, int, bool):
                    continue
                if type(cell) is float:
                    if not math.isfinite(cell):
                        raise ValueError("Numeric cell values must be finite")
                    continue
                raise ValueError(f"Unsupported cell type: {type(cell).__name__}")


def _write_validated_matrix_sheet(workbook: Workbook, matrix: WorksheetMatrix) -> Worksheet:
    worksheet = workbook.create_sheet(title=matrix.sheet_name)
    for column, value in enumerate(matrix.headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=value)
        cell.data_type = "s"

    for row_number, row_values in enumerate(matrix.rows, start=2):
        for column, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
    return worksheet


def write_matrix_sheet(workbook: Workbook, matrix: WorksheetMatrix) -> Worksheet:
    """Append one validated matrix as a worksheet to an existing workbook."""
    _validate_matrices([matrix])
    return _write_validated_matrix_sheet(workbook, matrix)


def build_matrix_workbook(matrices: list[WorksheetMatrix]) -> bytes:
    """Build one workbook after validating the complete payload and return XLSX bytes."""
    _validate_matrices(matrices)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for matrix in matrices:
        _write_validated_matrix_sheet(workbook, matrix)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def validate_export_filename(filename: str) -> str:
    """Validate path safety and normalize a requested download name to ASCII."""
    if not isinstance(filename, str) or not filename:
        raise ValueError("Export filename must not be empty")
    if "\r" in filename or "\n" in filename:
        raise ValueError("Export filename cannot contain line breaks")
    if "/" in filename or "\\" in filename or ":" in filename:
        raise ValueError("Export filename must not contain path components")
    if len(filename) < 5 or filename[-5:].lower() != ".xlsx":
        raise ValueError("Export filename must end with .xlsx")

    stem = filename[:-5]
    safe_stem = _INVALID_FILENAME_CHARS.sub("_", stem).strip(" ._")
    if safe_stem in ("", ".", ".."):
        safe_stem = "export"
    return f"{safe_stem}.xlsx"
