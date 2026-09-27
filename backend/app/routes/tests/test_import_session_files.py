from __future__ import annotations

import io
import multiprocessing
import os
import threading
import uuid
from hashlib import sha256

import openpyxl
import pytest

import app.services.import_parsers.v1_standard  # noqa: F401
from app.db.models import ImportSession
from app.services.import_sessions import (
    ImportSessionError,
    read_import_workbook,
)
from app.storage.dependencies import get_object_storage
from app.storage.protocol import StoredObject
from tests.helpers import auth_headers, create_soldier


class FakeStorage:
    def __init__(self, objects=None):
        self.objects = objects or {}
        self.put_count = 0

    def put_bytes(self, *, key, data, content_type, sha256):
        self.objects[key] = data
        self.put_count += 1
        return StoredObject(key, len(data), sha256, None, None)

    def open_read(self, *, key):
        return io.BytesIO(self.objects[key]), len(self.objects[key])


def _workbook_bytes():
    workbook = openpyxl.Workbook()
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def test_import_workbook_reads_verified_storage_and_legacy_fallback():
    content = _workbook_bytes()
    stored = ImportSession(filename="a.xlsx", raw_excel=None)
    stored.id = uuid.uuid4()
    key = f"import_workbook/{stored.id}"
    storage = FakeStorage({key: content})
    stored.storage_key = key
    stored.storage_sha256 = sha256(content).hexdigest()
    stored.storage_size = len(content)
    assert read_import_workbook(None, stored, storage) == content

    legacy = ImportSession(filename="legacy.xlsx", raw_excel=content)
    assert read_import_workbook(None, legacy, storage) == content


def test_import_storage_outage_does_not_fall_back_to_legacy_bytes():
    content = _workbook_bytes()
    stored = ImportSession(filename="a.xlsx", raw_excel=content)
    stored.id = uuid.uuid4()
    stored.storage_key = f"import_workbook/{uuid.uuid4()}"
    stored.storage_sha256 = sha256(content).hexdigest()
    stored.storage_size = len(content)
    with pytest.raises(ImportSessionError, match="storage_unavailable"):
        read_import_workbook(None, stored, FakeStorage({}))


def test_xlsx_validator_rejects_non_zip_before_parser():
    from app.services.file_validation import validate_xlsx

    with pytest.raises(ValueError, match="invalid_file_type"):
        validate_xlsx(b"PK\x03\x04bad")


@pytest.mark.skipif(os.name == "nt", reason="bounded parser execution requires the Linux container path")
def test_import_session_create_reparse_and_confirm_use_object_storage(admin_session):
    from app.services.import_sessions import confirm_session, create_session, reparse_session
    from tests.helpers import create_soldier

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "soldiers"
    sheet.append(["personal_number", "full_name", "rank", "gender", "is_officer", "hierarchy_node_name", "enrolled_at", "enlistment_date", "phone", "email"])
    sheet.append(["9123456", "Storage Import", "", "", "", "", "", "", "", ""])
    content = _workbook_bytes_from(workbook)
    actor = create_soldier(admin_session, personal_number=f"adm{uuid.uuid4().hex[:8]}", role="admin")
    storage = FakeStorage()

    imported = create_session(
        admin_session, filename="users.xlsx", content=content, actor=actor, parser_id="v1_standard", storage=storage,
    )
    assert imported.raw_excel is None
    assert imported.storage_key == f"import_workbook/{imported.id}"
    assert read_import_workbook(admin_session, imported, storage) == content
    reparse_session(admin_session, session_id=imported.id, actor=actor, storage=storage)
    confirm_session(admin_session, session_id=imported.id, actor=actor, storage=storage)
    admin_session.commit()
    assert imported.raw_excel is None
    assert storage.put_count == 1


def _workbook_bytes_from(workbook):
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def test_import_upload_rejects_invalid_xlsx_before_put(client, admin_session):
    storage = FakeStorage()
    client.app.dependency_overrides[get_object_storage] = lambda: storage
    admin = create_soldier(admin_session, personal_number=f"adm{uuid.uuid4().hex[:8]}", role="admin")
    try:
        response = client.post(
            "/api/import/sessions", files={"file": ("broken.xlsx", b"PK\x03\x04malformed", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=auth_headers(admin),
        )
        assert response.status_code == 400
        assert response.json()["detail"] == "invalid_file_type"
        assert not storage.objects
    finally:
        client.app.dependency_overrides.pop(get_object_storage, None)

def test_parser_does_not_run_when_memory_limit_cannot_be_applied(monkeypatch):
    from types import SimpleNamespace

    from app.services import import_sessions as import_service
    from app.services.import_sessions import _parse_workbook_child

    monkeypatch.setattr(import_service.os, "name", "posix")
    attempts = []
    monkeypatch.setitem(__import__("sys").modules, "resource", SimpleNamespace(setrlimit=lambda *_args: (_ for _ in ()).throw(OSError("denied"))))
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_args, **_kwargs: attempts.append("parsed"))

    class Connection:
        messages = None

        def send(self, message):
            self.messages = message

        def close(self):
            pass

    connection = Connection()
    _parse_workbook_child(connection, _workbook_bytes(), "v1_standard")

    assert connection.messages[0] is False
    assert attempts == []


def test_parser_result_larger_than_pipe_capacity_is_drained_without_deadlock():
    from app.services.import_sessions import _receive_parser_message

    reader, writer = multiprocessing.get_context().Pipe(duplex=False)
    payload = b"x" * (2 * 1024 * 1024)
    sender = threading.Thread(target=lambda: writer.send((True, "v1_standard", payload)))
    sender.start()
    try:
        message = _receive_parser_message(reader, timeout_seconds=5)
        sender.join(timeout=5)
        assert not sender.is_alive()
        assert message == (True, "v1_standard", payload)
    finally:
        reader.close()
        writer.close()
        sender.join(timeout=5)


def test_import_upload_rejects_mismatched_mime_before_parse_or_storage(client, admin_session, monkeypatch):
    from app.routes import import_sessions as import_routes

    storage = FakeStorage()
    client.app.dependency_overrides[get_object_storage] = lambda: storage
    admin = create_soldier(admin_session, personal_number=f"adm{uuid.uuid4().hex[:8]}", role="admin")
    parse_calls = []
    def record_parse(*args, **kwargs):
        from types import SimpleNamespace

        parse_calls.append("called")
        return SimpleNamespace(id=uuid.uuid4(), storage_key=None, parsed_state=None)

    monkeypatch.setattr(import_routes, "create_session", record_parse)
    try:
        response = client.post(
            "/api/import/sessions",
            files={"file": ("valid.xlsx", _workbook_bytes(), "application/zip")},
            headers=auth_headers(admin),
        )
        assert response.status_code == 400
        assert response.json()["detail"] == "invalid_content_type"
        assert storage.put_count == 0
        assert parse_calls == []
        assert admin_session.query(ImportSession).count() == 0
    finally:
        client.app.dependency_overrides.pop(get_object_storage, None)
