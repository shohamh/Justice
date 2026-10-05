"""Authenticated endpoints that create bounded XLSX exports."""

from __future__ import annotations

import io

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Request, status
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, ValidationError
from sqlalchemy.orm import Session

from app.auth.deps import require_duty_manager_or_admin, require_password_changed
from app.db.models import Soldier
from app.db.session import get_session
from app.routes import config_export
from app.routes.import_excel import EXPORT_DATA_SHEETS, _write_export_data_sheets
from app.services.excel_bilingual import finalize_bilingual_workbook
from app.services.xlsx_export import (
    MAX_EXPORT_CELLS,
    MAX_EXPORT_REQUEST_BYTES,
    WorksheetMatrix,
    build_matrix_workbook,
    validate_export_filename,
    write_matrix_sheet,
)

router = APIRouter()

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class MatrixExportRequest(BaseModel):
    """Strict request envelope for a client-prepared worksheet matrix."""

    model_config = ConfigDict(strict=True, extra="forbid")

    filename: str
    matrix: WorksheetMatrix


class PlanningExportRequest(BaseModel):
    """Strict request envelope for selected planning matrices and server sheets."""

    model_config = ConfigDict(strict=True, extra="forbid")

    filename: str
    tables: list[WorksheetMatrix]
    config_sheets: list[str]
    data_sheets: list[str]


async def _read_bounded_body(request: Request) -> bytes:
    body = bytearray()
    async for chunk in request.stream():
        if len(body) + len(chunk) > MAX_EXPORT_REQUEST_BYTES:
            raise HTTPException(
                status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                detail="export_request_too_large",
            )
        body.extend(chunk)
    return bytes(body)


@router.post("/exports/xlsx")
async def export_xlsx_matrix(
    request: Request,
    _user: Soldier = Depends(require_password_changed),
) -> Response:
    """Build an XLSX attachment from a bounded matrix supplied by the client."""
    raw_body = await _read_bounded_body(request)
    try:
        payload = MatrixExportRequest.model_validate_json(raw_body, strict=True)
    except (ValidationError, ValueError) as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="invalid_export_request",
        ) from exc

    try:
        filename = validate_export_filename(payload.filename)
        workbook = build_matrix_workbook([payload.matrix])
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="invalid_export_request",
        ) from exc

    return Response(
        content=workbook,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/exports/planning-xlsx")
async def export_planning_xlsx(
    request: Request,
    session: Session = Depends(get_session),
    actor: Soldier = Depends(require_duty_manager_or_admin),
) -> Response:
    """Combine client-prepared planning tables with authorized server exports."""
    raw_body = await _read_bounded_body(request)
    try:
        payload = PlanningExportRequest.model_validate_json(raw_body, strict=True)
    except (ValidationError, ValueError) as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="invalid_export_request",
        ) from exc

    try:
        filename = validate_export_filename(payload.filename)
        matrix_names = [matrix.sheet_name for matrix in payload.tables]
        if any(name not in {"transparency", "sub_units"} for name in matrix_names):
            raise ValueError("Unknown client planning table")
        if len(matrix_names) != len(set(matrix_names)):
            raise ValueError("Duplicate client planning table")
        if len(payload.config_sheets) != len(set(payload.config_sheets)):
            raise ValueError("Duplicate config sheet")
        if len(payload.data_sheets) != len(set(payload.data_sheets)):
            raise ValueError("Duplicate data sheet")
        if any(name not in config_export._WRITERS for name in payload.config_sheets):
            raise ValueError("Unknown config sheet")
        if any(name not in EXPORT_DATA_SHEETS for name in payload.data_sheets):
            raise ValueError("Unknown data sheet")

        client_cell_count = sum(
            len(matrix.headers) + len(matrix.rows) * len(matrix.headers)
            for matrix in payload.tables
        )
        if client_cell_count > MAX_EXPORT_CELLS:
            raise ValueError("Planning export cell limit exceeded")
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="invalid_export_request",
        ) from exc

    config_sheets = config_export._filter_export_sheets(list(payload.config_sheets), actor)

    ordered_matrices = sorted(
        payload.tables,
        key=lambda matrix: ("transparency", "sub_units").index(matrix.sheet_name),
    )
    if not (ordered_matrices or config_sheets or payload.data_sheets):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="no_exportable_sheets",
        )

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    try:
        for matrix in ordered_matrices:
            write_matrix_sheet(workbook, matrix)
        for sheet_name in config_sheets:
            config_export._WRITERS[sheet_name](workbook, session)
        _write_export_data_sheets(workbook, session, set(payload.data_sheets))
        finalize_bilingual_workbook(workbook)
        buffer = io.BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="invalid_export_request",
        ) from exc
    finally:
        workbook.close()

    return Response(
        content=content,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
